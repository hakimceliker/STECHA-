"""Document Processing Service - Extract text from PDFs, DOCx, XLSX files"""

import os
import tempfile
from typing import Optional
from pathlib import Path

try:
    import PyPDF2
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


class DocumentProcessor:
    """Process various document formats and extract text content"""

    SUPPORTED_FORMATS = {".pdf", ".docx", ".xlsx", ".txt"}
    MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB

    @staticmethod
    def validate_file(file_path: str, file_size: int) -> tuple[bool, str]:
        """Validate file format and size"""
        if file_size > DocumentProcessor.MAX_FILE_SIZE:
            return False, f"File exceeds maximum size of 50 MB"

        ext = Path(file_path).suffix.lower()
        if ext not in DocumentProcessor.SUPPORTED_FORMATS:
            return False, f"Unsupported format. Supported: {', '.join(DocumentProcessor.SUPPORTED_FORMATS)}"

        return True, "OK"

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF file"""
        if not HAS_PYPDF:
            raise ImportError("PyPDF2 not installed. Install with: pip install PyPDF2")

        text = []
        try:
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    text.append(page.extract_text())
        except Exception as e:
            raise ValueError(f"Error reading PDF: {str(e)}")

        return "\n".join(text)

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX file"""
        if not HAS_DOCX:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")

        text = []
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    text.append(para.text)
        except Exception as e:
            raise ValueError(f"Error reading DOCX: {str(e)}")

        return "\n".join(text)

    @staticmethod
    def extract_text_from_xlsx(file_path: str) -> str:
        """Extract text from XLSX file"""
        if not HAS_XLSX:
            raise ImportError("openpyxl not installed. Install with: pip install openpyxl")

        text = []
        try:
            workbook = openpyxl.load_workbook(file_path)
            for sheet in workbook.sheetnames:
                worksheet = workbook[sheet]
                text.append(f"\n=== Sheet: {sheet} ===\n")
                for row in worksheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        text.append(row_text)
        except Exception as e:
            raise ValueError(f"Error reading XLSX: {str(e)}")

        return "\n".join(text)

    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                raise ValueError(f"Error reading TXT: {str(e)}")

    @staticmethod
    def extract_text(file_path: str) -> str:
        """Extract text from any supported document format"""
        ext = Path(file_path).suffix.lower()

        if ext == ".pdf":
            return DocumentProcessor.extract_text_from_pdf(file_path)
        elif ext == ".docx":
            return DocumentProcessor.extract_text_from_docx(file_path)
        elif ext == ".xlsx":
            return DocumentProcessor.extract_text_from_xlsx(file_path)
        elif ext == ".txt":
            return DocumentProcessor.extract_text_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported format: {ext}")


class DocumentSearcher:
    """Search within extracted document text"""

    @staticmethod
    def search(text: str, query: str, context_lines: int = 3) -> list[dict]:
        """Search for query in text and return results with context"""
        results = []
        lines = text.split('\n')
        query_lower = query.lower()

        for i, line in enumerate(lines):
            if query_lower in line.lower():
                start = max(0, i - context_lines)
                end = min(len(lines), i + context_lines + 1)
                context = '\n'.join(lines[start:end])

                results.append({
                    "line_number": i + 1,
                    "matched_line": line,
                    "context": context,
                    "position": line.lower().index(query_lower)
                })

        return results

    @staticmethod
    def summarize(text: str, max_sentences: int = 5) -> str:
        """Create a simple summary by extracting key sentences"""
        sentences = text.split('.')
        sentences = [s.strip() for s in sentences if s.strip()]

        if len(sentences) <= max_sentences:
            return '. '.join(sentences) + '.'

        scored_sentences = []
        for sentence in sentences:
            words = sentence.lower().split()
            score = sum(1 for word in words if len(word) > 3)
            scored_sentences.append((sentence, score))

        scored_sentences.sort(key=lambda x: x[1], reverse=True)
        summary = '. '.join(s[0] for s in scored_sentences[:max_sentences]) + '.'
        return summary
